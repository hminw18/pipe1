from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import Any, TypedDict

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from sewerpipe_inspector.defect_taxonomy import (
    GRADE_ORDER,
    condition_code as taxonomy_condition_code,
    defect_definition,
    defect_code as taxonomy_defect_code,
    defect_score,
    display_condition_item,
    display_defect_item,
)
from sewerpipe_inspector.state_grading import (
    BOUNDARY_CONDITION_ITEMS,
    compute_pipe_state_grades,
    format_state_distance,
    format_state_section_label,
    format_state_time_range,
    format_state_value,
)


GRADE_FILL = {
    "대": PatternFill(fill_type="solid", fgColor="FF7F7F"),
    "중": PatternFill(fill_type="solid", fgColor="FFA500"),
    "소": PatternFill(fill_type="solid", fgColor="90EE90"),
}

HEADER_FILL = PatternFill(fill_type="solid", fgColor="E8EEF7")
LABEL_FILL = PatternFill(fill_type="solid", fgColor="F2F2F2")
LIGHT_BLUE_FILL = PatternFill(fill_type="solid", fgColor="EAF3F8")
THIN_SIDE = Side(style="thin", color="808080")
THIN_BORDER = Border(
    left=THIN_SIDE,
    right=THIN_SIDE,
    top=THIN_SIDE,
    bottom=THIN_SIDE,
)

AGGREGATE_DEFECT_ITEMS = [
    "균열(길이)",
    "균열(원주)",
    "균열(복합)",
    "표면손상",
    "라이닝결함",
    "좌굴",
    "변형",
    "파손",
    "붕괴",
    "영구장애물",
    "천공",
    "연결관돌출",
    "연결관접합부",
    "이음부이탈",
    "이음부손상",
    "이음부단차",
    "역경사",
    "침하",
    "내피생성",
    "토사퇴적",
    "폐유부착",
    "임시장애물",
    "뿌리침입",
    "침입수",
    "막힘",
]

AGGREGATE_CONDITION_ITEMS = [
    "조사시작(순방향)",
    "조사시작(역방향)",
    "조사완료(순방향)",
    "조사완료(역방향)",
    "조사중단",
    "시야상실",
    "중심상실",
    "대상없음",
    "특이사항",
    "연결관미사용",
    "연결관존재함",
    "이음부(접합부)존재",
    "재질변경",
    "라이닝변화",
    "이상없음",
    "단위길이",
]

AGGREGATE_CONDITION_LABELS = {
    "조사시작(순방향)": "(상태)\n관로조사시작\n(순방향)",
    "조사시작(역방향)": "(상태)\n관로조사시작\n(역방향)",
    "조사완료(순방향)": "(상태)\n관로조사완료\n(순방향)",
    "조사완료(역방향)": "(상태)\n관로조사완료\n(역방향)",
    "조사중단": "(상태)\n관로조사중단",
    "시야상실": "(상태)\n시야상실",
    "중심상실": "(상태)\n중심상실",
    "대상없음": "(상태)\n대상없음",
    "특이사항": "(상태)\n특이사항",
    "연결관미사용": "(상태)\n연결관미사용",
    "연결관존재함": "(상태)\n연결관존재함",
    "이음부(접합부)존재": "(상태)\n이음부(접합부)존재",
    "재질변경": "(상태)\n재질변경",
    "라이닝변화": "(상태)\n라이닝변화",
    "이상없음": "(상태)\n이상없음",
    "단위길이": "(상태)\n단위길이",
}


class DefectReportRow(TypedDict):
    timestamp_ms: int
    drive_direction: str
    distance_m: float | None
    item_category: str | None
    condition_item: str | None
    defect_item: str | None
    grade: str | None
    quadrant: str | None
    manhole_defect_depth_m: float | None
    memo: str | None
    image_path: str


class ReportWorkbookData(TypedDict):
    context: Any
    report: Any
    pipe_info: Any
    upstream_manhole: Any
    downstream_manhole: Any
    actual_survey: Any
    defects: list[DefectReportRow]


def _format_hms(timestamp_ms: int) -> str:
    seconds = timestamp_ms // 1000
    hours = seconds // 3600
    minutes = (seconds % 3600) // 60
    secs = seconds % 60
    return f"{hours:02d}:{minutes:02d}:{secs:02d}"


def _value(row: Any, key: str, default: str = "") -> Any:
    if row is None:
        return default
    try:
        value = row[key]
    except Exception:
        value = getattr(row, key, default)
    return default if value is None else value


def _float_value(row: Any, key: str) -> float | None:
    value = _value(row, key, None)
    if value in (None, ""):
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def _number_or_blank(value: float | int | None) -> float | int | str:
    return "" if value is None else value


def _format_meters(value: float | int | None) -> str:
    if value is None:
        return ""
    return f"{float(value):,.2f}m"


def _state_grade_formula(score_cell: str) -> str:
    return (
        f'=IF({score_cell}>69,5,'
        f'IF({score_cell}>39,4,'
        f'IF({score_cell}>19,3,'
        f'IF({score_cell}>9,2,1))))'
    )


def _safe_sheet_title(title: str, used: set[str]) -> str:
    invalid = '[]:*?/\\'
    base = "".join("_" if ch in invalid else ch for ch in str(title).strip())
    base = (base or "보고서")[:31]
    candidate = base
    suffix = 1
    while candidate in used:
        marker = f"_{suffix}"
        candidate = f"{base[:31 - len(marker)]}{marker}"
        suffix += 1
    used.add(candidate)
    return candidate


def _defect_display_name(defect_item: str | None, grade: str | None) -> str:
    if not defect_item:
        return ""
    return f"{defect_item}({grade})" if grade else defect_item


def _defect_code(category: str | None, defect_item: str | None) -> str:
    return taxonomy_defect_code(category, defect_item)


def _condition_code(category: str | None, condition_item: str | None) -> str:
    return taxonomy_condition_code(category, condition_item)


def _apply_table_cell_style(cell, *, fill=None, bold: bool = False) -> None:
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = THIN_BORDER
    if fill is not None:
        cell.fill = fill
    if bold:
        cell.font = Font(bold=True)


def _style_range(
    ws,
    min_row: int,
    max_row: int,
    min_col: int,
    max_col: int,
    *,
    fill=None,
    bold: bool = False,
) -> None:
    for row in ws.iter_rows(
        min_row=min_row,
        max_row=max_row,
        min_col=min_col,
        max_col=max_col,
    ):
        for cell in row:
            _apply_table_cell_style(cell, fill=fill, bold=bold)


def _merge_and_style_outer_border(
    ws,
    min_row: int,
    max_row: int,
    min_col: int,
    max_col: int,
    *,
    fill=None,
    bold: bool = False,
) -> None:
    ws.merge_cells(
        start_row=min_row,
        start_column=min_col,
        end_row=max_row,
        end_column=max_col,
    )
    for row in range(min_row, max_row + 1):
        for col in range(min_col, max_col + 1):
            cell = ws.cell(row, col)
            cell.alignment = Alignment(
                horizontal="center", vertical="center", wrap_text=True
            )
            if fill is not None:
                cell.fill = fill
            if bold:
                cell.font = Font(bold=True)
            cell.border = Border(
                left=THIN_SIDE if col == min_col else None,
                right=THIN_SIDE if col == max_col else None,
                top=THIN_SIDE if row == min_row else None,
                bottom=THIN_SIDE if row == max_row else None,
            )


def _align_range(
    ws,
    min_row: int,
    max_row: int,
    min_col: int,
    max_col: int,
    *,
    bold: bool = False,
) -> None:
    for row in ws.iter_rows(
        min_row=min_row,
        max_row=max_row,
        min_col=min_col,
        max_col=max_col,
    ):
        for cell in row:
            cell.alignment = Alignment(
                horizontal="center", vertical="center", wrap_text=True
            )
            if bold:
                cell.font = Font(bold=True)


def _append_title(ws, title: str, row_index: int, width: int = 4) -> int:
    ws.cell(row=row_index, column=1, value=title)
    ws.merge_cells(
        start_row=row_index,
        start_column=1,
        end_row=row_index,
        end_column=width,
    )
    cell = ws.cell(row=row_index, column=1)
    cell.font = Font(bold=True, size=13)
    cell.fill = HEADER_FILL
    cell.alignment = Alignment(horizontal="center")
    return row_index + 1


def _append_pair_table(ws, row_index: int, pairs: list[tuple[str, Any]]) -> int:
    for idx in range(0, len(pairs), 2):
        left = pairs[idx]
        right = pairs[idx + 1] if idx + 1 < len(pairs) else ("", "")
        values = [left[0], left[1], right[0], right[1]]
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=row_index, column=col, value=value)
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            if col in (1, 3):
                cell.fill = LABEL_FILL
                cell.font = Font(bold=True)
        row_index += 1
    return row_index + 1


class ReportService:
    def generate_inspection_report(
        self,
        report_path: Path,
        context,
        report,
        pipe_info,
        upstream_manhole,
        downstream_manhole,
        actual_survey,
        defects: list[DefectReportRow],
    ) -> Path:
        wb = Workbook()
        self._enable_recalculation(wb)
        ws_basic = wb.active
        if ws_basic is None:
            raise RuntimeError("Workbook initialization failed")
        ws_basic.title = "기본정보"
        self._build_basic_sheet(ws_basic, context, report)

        ws_pipe = wb.create_sheet("관로조사정보")
        self._build_pipe_sheet(
            ws_pipe,
            pipe_info,
            upstream_manhole,
            downstream_manhole,
            actual_survey,
        )

        ws_defects = wb.create_sheet("결함목록")
        self._build_defect_sheet(ws_defects, defects)

        report_path.parent.mkdir(parents=True, exist_ok=True)
        wb.save(report_path)
        return report_path

    def generate_internal_defect_workbook(
        self,
        report_path: Path,
        business_context: Any,
        reports: list[ReportWorkbookData],
        include_photos: bool = False,
    ) -> Path:
        wb = Workbook()
        self._enable_recalculation(wb)
        ws_summary = wb.active
        if ws_summary is None:
            raise RuntimeError("Workbook initialization failed")
        ws_summary.title = "내부조사(CCTV)집계표"
        self._build_internal_summary_sheet(ws_summary, reports)

        used_titles = set(wb.sheetnames)
        for item in reports:
            title = _safe_sheet_title(
                str(_value(item["report"], "report_number", "보고서")),
                used_titles,
            )
            self._build_internal_defect_sheet(
                wb.create_sheet(title), item, include_photos=include_photos
            )

        report_path.parent.mkdir(parents=True, exist_ok=True)
        wb.save(report_path)
        return report_path

    def generate_defect_aggregate_workbook(
        self,
        report_path: Path,
        business_context: Any,
        reports: list[ReportWorkbookData],
    ) -> Path:
        wb = Workbook()
        self._enable_recalculation(wb)
        ws = wb.active
        if ws is None:
            raise RuntimeError("Workbook initialization failed")
        ws.title = "Sheet1"
        self._build_defect_aggregate_sheet(ws, business_context, reports)
        report_path.parent.mkdir(parents=True, exist_ok=True)
        wb.save(report_path)
        return report_path

    @staticmethod
    def _enable_recalculation(wb: Workbook) -> None:
        try:
            wb.calculation.fullCalcOnLoad = True
            wb.calculation.forceFullCalc = True
        except Exception:
            return

    def _build_internal_defect_sheet(
        self, ws, data: ReportWorkbookData, *, include_photos: bool = False
    ) -> None:
        report = data["report"]
        pipe_info = data["pipe_info"]
        upstream = data["upstream_manhole"]
        downstream = data["downstream_manhole"]
        actual = data["actual_survey"]
        defects = [
            defect
            for defect in data["defects"]
            if (defect.get("item_category") or "관로") == "관로"
        ]
        display_rows = self._internal_defect_display_rows(defects)
        data_start = 9
        last_data_row = max(data_start, data_start + len(display_rows) - 1)
        summary_row = last_data_row + 1

        self._build_internal_defect_header(
            ws, summary_row, include_photos=include_photos
        )
        self._fill_internal_defect_metadata(
            ws, report, pipe_info, upstream, downstream
        )

        for offset, row_data in enumerate(display_rows):
            row_index = data_start + offset
            values = [
                row_data["section_label"],
                row_data["position"],
                row_data["kind"],
                row_data["item"],
                row_data["code"],
                row_data["structural_score"],
                row_data["operational_score"],
                row_data["memo"],
                "",
                "",
                "",
                "",
            ]
            for col, value in enumerate(values, start=1):
                _apply_table_cell_style(ws.cell(row_index, col, value))
            if include_photos:
                _merge_and_style_outer_border(ws, row_index, row_index, 13, 22)
                image_path = Path(str(row_data.get("image_path") or ""))
                if image_path.exists():
                    image = XLImage(str(image_path))
                    image.width = 120
                    image.height = 68
                    ws.add_image(image, f"M{row_index}")
                    ws.row_dimensions[row_index].height = 54

        for first_row, last_row in self._internal_unit_ranges(display_rows, data_start):
            ws.cell(first_row, 9, f"=MAX(F{first_row}:F{last_row})")
            ws.cell(first_row, 10, _state_grade_formula(f"I{first_row}"))
            ws.cell(first_row, 11, f"=MAX(G{first_row}:G{last_row})")
            ws.cell(first_row, 12, _state_grade_formula(f"K{first_row}"))

        ws.merge_cells(start_row=summary_row, start_column=1, end_row=summary_row, end_column=8)
        ws.merge_cells(start_row=summary_row, start_column=9, end_row=summary_row, end_column=10)
        ws.merge_cells(start_row=summary_row, start_column=11, end_row=summary_row, end_column=12)
        ws.cell(summary_row, 1, "'조사구간' 구조적, 운영적 상태등급")
        ws.cell(
            summary_row,
            9,
            f'=IFERROR(ROUNDUP(SUM(J{data_start}:J{last_data_row})/'
            f'COUNTA(J{data_start}:J{last_data_row}),2),"오류")',
        )
        ws.cell(
            summary_row,
            11,
            f'=IFERROR(ROUNDUP(SUM(L{data_start}:L{last_data_row})/'
            f'COUNTA(L{data_start}:L{last_data_row}),2),"오류")',
        )
        _style_range(ws, summary_row, summary_row, 1, 12, fill=LABEL_FILL, bold=True)
        ws.row_dimensions[summary_row].height = 16.5

    def _build_internal_defect_header(
        self,
        ws,
        summary_row: int | None,
        *,
        include_photos: bool = False,
    ) -> None:
        for col in range(1, 13):
            ws.column_dimensions[get_column_letter(col)].width = (
                6.625 if col in (1, 5, 9) else 13
            )
        for col in range(13, 23):
            ws.column_dimensions[get_column_letter(col)].width = 4.5
        ws.column_dimensions["D"].width = 16.125
        ws.column_dimensions["H"].width = 10
        for row in range(2, 9):
            ws.row_dimensions[row].height = 16.5

        merges = [
            "A1:L1",
            "A2:B2",
            "C2:D2",
            "F2:G2",
            "K2:L2",
            "A3:B3",
            "C3:D3",
            "E3:G3",
            "H3:I3",
            "K3:L3",
            "M3:Q3",
            "R3:V3",
            "A4:D4",
            "E4:G4",
            "H4:J4",
            "K4:L4",
            "A5:D5",
            "E5:G5",
            "H5:J5",
            "K5:L5",
            "A6:A8",
            "B6:B8",
            "C6:H6",
            "I6:L6",
            "C7:C8",
            "D7:D8",
            "E7:E8",
            "F7:F8",
            "G7:G8",
            "H7:H8",
            "I7:J7",
            "K7:L7",
        ]
        for merge_range in merges:
            ws.merge_cells(merge_range)

        ws.cell(1, 1, "<내부결함 판독표>")
        ws.cell(2, 1, "보고서 번호")
        ws.cell(2, 5, "관종")
        ws.cell(2, 8, "규격")
        ws.cell(2, 10, "배수방식")
        ws.cell(3, 1, "시작맨홀")
        ws.cell(3, 5, "끝맨홀")
        ws.cell(3, 10, "주행거리")
        ws.cell(3, 13, "구조")
        ws.cell(3, 18, "운영")
        ws.cell(4, 1, "조사구간 구조적 상태등급")
        ws.cell(4, 8, "조사구간 운영적 상태등급")
        ws.cell(5, 1, "단위구간 구조적 최고등급")
        ws.cell(5, 8, "단위구간 운영적 최고등급")
        ws.cell(6, 1, "단위\n구간")
        ws.cell(6, 2, "위치")
        ws.cell(6, 3, "하수관로 결함")
        ws.cell(6, 9, "단위구간 상태등급")
        if include_photos:
            ws.merge_cells("M6:V8")
            ws.cell(6, 13, "사진")
        for index, grade in enumerate((1, 2, 3, 4, 5), start=13):
            ws.cell(4, index, grade)
            ws.cell(5, index, f'=COUNTIFS($J9:$J2000,"={grade}")')
        for index, grade in enumerate((1, 2, 3, 4, 5), start=18):
            ws.cell(4, index, grade)
            ws.cell(5, index, f'=COUNTIFS($L9:$L2000,"={grade}")')
        for col, value in enumerate(
            ["구분", "결함", "코드", "구조\n점수", "운영\n점수", "비고"],
            start=3,
        ):
            ws.cell(7, col, value)
        ws.cell(7, 9, "구조")
        ws.cell(7, 11, "운영")
        ws.cell(8, 9, "점수")
        ws.cell(8, 10, "등급")
        ws.cell(8, 11, "점수")
        ws.cell(8, 12, "등급")

        if summary_row is not None:
            ws.cell(4, 5, f"=I{summary_row}")
            ws.cell(4, 11, f"=K{summary_row}")
            ws.cell(5, 5, f"=MAX(J9:J{summary_row - 1})")
            ws.cell(5, 11, f"=MAX(L9:L{summary_row - 1})")

        _align_range(ws, 1, 1, 1, 12, bold=True)
        _style_range(ws, 2, 8, 1, 12, fill=LIGHT_BLUE_FILL, bold=True)
        _style_range(ws, 3, 5, 13, 22, bold=True)
        if include_photos:
            _style_range(ws, 6, 8, 13, 22, bold=True)
        ws.cell(1, 1).alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=False
        )
        ws.cell(1, 1).font = Font(bold=True, size=13)

    def _fill_internal_defect_metadata(
        self, ws, report, pipe_info, upstream, downstream
    ) -> None:
        ws.cell(2, 3, _value(report, "report_number"))
        ws.cell(2, 6, _value(report, "pipe_type"))
        ws.cell(2, 9, _value(report, "specification"))
        ws.cell(2, 11, _value(report, "drain_type"))
        ws.cell(3, 3, _value(upstream, "manhole_number"))
        ws.cell(3, 8, _value(downstream, "manhole_number"))
        ws.cell(3, 11, _number_or_blank(_float_value(pipe_info, "total_drive_distance_m")))

    def _internal_defect_display_rows(
        self, defects: list[DefectReportRow]
    ) -> list[dict[str, object]]:
        rows: list[dict[str, object]] = []
        for drive_direction in ("순주행", "역주행"):
            direction_rows = [
                row for row in defects if row.get("drive_direction") == drive_direction
            ]
            reverse = drive_direction == "역주행"
            direction_rows.sort(
                key=lambda row: (
                    -float(row.get("distance_m") or 0)
                    if reverse
                    else float(row.get("distance_m") or 0),
                    int(row.get("timestamp_ms") or 0),
                )
            )
            section_index = 1
            for row_index, defect in enumerate(direction_rows):
                condition_item = defect.get("condition_item") or ""
                defect_item = defect.get("defect_item") or ""
                grade = defect.get("grade") or ""
                distance = float(defect.get("distance_m") or 0)
                display_distance = -distance if reverse and distance != 0 else distance
                is_defect = bool(defect_item)
                definition = (
                    defect_definition("관로", defect_item) if is_defect else None
                )
                score = (
                    defect_score("관로", defect_item, grade)
                    if is_defect
                    else None
                )
                kind = definition.defect_type if definition is not None else "상태"
                section_label = (
                    f"{section_index}(역)" if reverse else str(section_index)
                )
                is_section_start = not rows or (
                    rows[-1].get("section_label") != section_label
                )
                rows.append(
                    {
                        "section_label": section_label,
                        "position": round(display_distance, 2),
                        "kind": kind,
                        "item": (
                            _defect_display_name(defect_item, grade)
                            if is_defect
                            else condition_item
                        ),
                        "code": (
                            _defect_code("관로", defect_item)
                            if is_defect
                            else _condition_code("관로", condition_item)
                        ),
                        "structural_score": (
                            score if kind == "구조" and score is not None else ""
                        ),
                        "operational_score": (
                            score if kind == "운영" and score is not None else ""
                        ),
                        "memo": (
                            "배경화면"
                            if is_section_start and condition_item.startswith("조사시작")
                            else defect.get("memo") or ""
                        ),
                        "image_path": defect.get("image_path") or "",
                    }
                )
                has_next = row_index < len(direction_rows) - 1
                if (
                    has_next
                    and condition_item in BOUNDARY_CONDITION_ITEMS
                    and not condition_item.startswith("조사시작")
                    and not is_defect
                ):
                    section_index += 1
        return rows

    @staticmethod
    def _internal_unit_ranges(
        display_rows: list[dict[str, object]], data_start: int
    ) -> list[tuple[int, int]]:
        if not display_rows:
            return [(data_start, data_start)]
        ranges: list[tuple[int, int]] = []
        start = data_start
        current = display_rows[0]["section_label"]
        for index, row in enumerate(display_rows, start=data_start):
            if row["section_label"] == current:
                continue
            ranges.append((start, index - 1))
            start = index
            current = row["section_label"]
        ranges.append((start, data_start + len(display_rows) - 1))
        return ranges

    def _build_internal_summary_sheet(
        self, ws, reports: list[ReportWorkbookData]
    ) -> None:
        ws.merge_cells("C2:AA2")
        ws.cell(2, 3, "하수관거 단위구간 보고서")
        headers = [
            "순서",
            "보고서번호",
            "관로번호",
            "시점맨홀번호",
            "종점맨홀번호",
            "관종",
            "관경",
            "관로연장",
            "조사연장",
            "조사율",
            "구조적",
            "운영적",
            "부분/전체",
            "구조적",
            "운영적",
            "정비판정",
        ]
        group_headers = [
            ("조사구간 상태등급", 12, 14),
            ("단위구간 최고등급", 15, 17),
            ("단위구간별 구조적 등급", 18, 22),
            ("단위구간별 운영적 등급", 23, 27),
        ]
        for title, start_col, end_col in group_headers:
            ws.merge_cells(start_row=3, start_column=start_col, end_row=3, end_column=end_col)
            ws.cell(3, start_col, title)
        for col, header in enumerate(headers, start=2):
            ws.cell(4, col, header)
        for offset, grade in enumerate((1, 2, 3, 4, 5), start=18):
            ws.cell(4, offset, grade)
        for offset, grade in enumerate((1, 2, 3, 4, 5), start=23):
            ws.cell(4, offset, grade)
        ws.cell(3, 28, "단위구간(개소)")
        ws.cell(4, 28, "단위구간(개소)")

        for col in range(2, 29):
            ws.column_dimensions[get_column_letter(col)].width = 13
        ws.column_dimensions["B"].width = 9
        ws.column_dimensions["E"].width = 16
        ws.column_dimensions["F"].width = 16
        ws.column_dimensions["H"].width = 8
        ws.column_dimensions["I"].width = 10
        ws.column_dimensions["J"].width = 10
        ws.column_dimensions["K"].width = 8
        ws.column_dimensions["AB"].width = 16
        ws.row_dimensions[2].height = 38.25
        ws.row_dimensions[3].height = 17.45
        ws.row_dimensions[4].height = 18.75

        _style_range(ws, 2, 4, 2, 28, fill=LIGHT_BLUE_FILL, bold=True)
        ws.cell(2, 3).font = Font(bold=True, size=14)

        data_start = 5
        for offset, item in enumerate(reports):
            row = data_start + offset
            report = item["report"]
            pipe_info = item["pipe_info"]
            upstream = item["upstream_manhole"]
            downstream = item["downstream_manhole"]
            defects = item["defects"]
            summary = compute_pipe_state_grades(defects)
            length_m = _float_value(pipe_info, "length_m")
            drive_m = _float_value(pipe_info, "total_drive_distance_m")
            structural_max = (
                max(section.structural_grade for section in summary.sections)
                if summary.sections
                else None
            )
            operational_max = (
                max(section.operational_grade for section in summary.sections)
                if summary.sections
                else None
            )
            structural_counts = {
                grade: sum(
                    1 for section in summary.sections if section.structural_grade == grade
                )
                for grade in range(1, 6)
            }
            operational_counts = {
                grade: sum(
                    1 for section in summary.sections if section.operational_grade == grade
                )
                for grade in range(1, 6)
            }
            values = [
                offset + 1,
                _value(report, "report_number"),
                _value(report, "pipe_number"),
                _value(upstream, "manhole_number"),
                _value(downstream, "manhole_number"),
                _value(report, "pipe_type"),
                _value(report, "specification"),
                _number_or_blank(length_m),
                _number_or_blank(drive_m),
                "" if not length_m else (drive_m or 0) / length_m,
                _number_or_blank(summary.structural_grade),
                _number_or_blank(summary.operational_grade),
                "전체" if int(_value(pipe_info, "is_completed", 0) or 0) else "부분",
                _number_or_blank(structural_max),
                _number_or_blank(operational_max),
                (
                    "정비대상"
                    if (structural_max or 0) >= 4 or (operational_max or 0) >= 4
                    else ""
                ),
            ]
            values.extend(structural_counts[grade] for grade in range(1, 6))
            values.extend(operational_counts[grade] for grade in range(1, 6))
            values.append(len(summary.sections))
            for col, value in enumerate(values, start=2):
                cell = ws.cell(row, col, value)
                _apply_table_cell_style(cell)
            ws.row_dimensions[row].height = 15.75

        total_row = data_start + len(reports)
        ws.cell(total_row, 3, "합계")
        ws.cell(total_row, 9, f"=SUM(I{data_start}:I{total_row - 1})")
        ws.cell(total_row, 10, f"=SUM(J{data_start}:J{total_row - 1})")
        ws.cell(total_row, 11, f"=IFERROR(J{total_row}/I{total_row},\"\")")
        _style_range(ws, total_row, total_row, 2, 28, fill=LABEL_FILL, bold=True)

    def _build_defect_aggregate_sheet(
        self, ws, business_context: Any, reports: list[ReportWorkbookData]
    ) -> None:
        aggregate_start_col = 13
        defect_start_col = aggregate_start_col + len(GRADE_ORDER)
        condition_start_col = defect_start_col + (len(AGGREGATE_DEFECT_ITEMS) * 3)
        location_start_col = condition_start_col + (len(AGGREGATE_CONDITION_ITEMS) * 3)
        max_col = location_start_col + 3
        max_letter = get_column_letter(max_col)
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max_col)
        ws.cell(1, 1, "하수관거 조사 집계표")
        ws.cell(3, 1, "사      업      명:")
        ws.merge_cells(f"C3:{max_letter}3")
        ws.cell(3, 3, _value(business_context, "business_name"))
        ws.cell(4, 1, "보고서번호범위:")
        ws.merge_cells(f"C4:{max_letter}4")
        ws.cell(4, 3, self._report_number_range(reports))
        ws.cell(5, 1, "보고서검색조건:")
        ws.merge_cells(f"C5:{max_letter}5")
        ws.cell(7, 1, "검   색   결   과:")
        ws.merge_cells("C7:K7")
        ws.cell(7, 3, f"{len(reports)}건")
        length_total = sum(_float_value(item["pipe_info"], "length_m") or 0 for item in reports)
        drive_total = sum(
            _float_value(item["pipe_info"], "total_drive_distance_m") or 0
            for item in reports
        )
        ws.cell(8, 1, "연장거리   합계:")
        ws.merge_cells("C8:K8")
        ws.cell(8, 3, _format_meters(length_total))
        ws.cell(9, 1, "주행거리   합계:")
        ws.merge_cells("C9:K9")
        ws.cell(9, 3, _format_meters(drive_total))
        ws.cell(10, 1, "미주행거리합계:")
        ws.merge_cells("C10:K10")
        ws.cell(10, 3, _format_meters(max(0.0, length_total - drive_total)))

        fixed_headers = [
            "관로번호",
            "상류\n맨홀",
            "하류\n맨홀",
            "배수\n방식",
            "관종",
            "규격",
            "연장",
            "주행\n거리",
            "(관로)\n구조적\n내부결함",
            "(관로)\n운영적\n내부결함",
        ]
        ws.merge_cells(start_row=11, start_column=1, end_row=11, end_column=2)
        ws.cell(11, 1, "보고서번호")
        for col, header in enumerate(fixed_headers, start=3):
            ws.cell(11, col, header)
        self._build_aggregate_group_headers(
            ws,
            aggregate_start_col,
            defect_start_col,
            condition_start_col,
            location_start_col,
        )

        data_start = 12
        for offset, item in enumerate(reports):
            row = data_start + offset
            self._fill_aggregate_report_row(ws, row, item)

        subtotal_row = data_start + len(reports)
        total_row = subtotal_row + 1
        ws.cell(subtotal_row, 1, "소계")
        ws.merge_cells(
            start_row=subtotal_row,
            start_column=1,
            end_row=subtotal_row,
            end_column=aggregate_start_col - 1,
        )
        ws.cell(total_row, 1, "합계")
        ws.merge_cells(
            start_row=total_row,
            start_column=1,
            end_row=total_row,
            end_column=aggregate_start_col - 1,
        )
        for col in range(aggregate_start_col, location_start_col):
            letter = get_column_letter(col)
            ws.cell(subtotal_row, col, f"=SUM({letter}{data_start}:{letter}{subtotal_row - 1})")
        for start_col in range(aggregate_start_col, location_start_col, 3):
            end_col = start_col + 2
            start_letter = get_column_letter(start_col)
            end_letter = get_column_letter(end_col)
            ws.merge_cells(
                start_row=total_row,
                start_column=start_col,
                end_row=total_row,
                end_column=end_col,
            )
            ws.cell(total_row, start_col, f"=SUM({start_letter}{subtotal_row}:{end_letter}{subtotal_row})")

        for col in range(1, max_col + 1):
            width = 13 if col < defect_start_col else 4.5
            ws.column_dimensions[get_column_letter(col)].width = width
        for col in (4, 8, 11, 12):
            ws.column_dimensions[get_column_letter(col)].width = 8
        ws.column_dimensions["A"].width = 18
        for col in range(condition_start_col, location_start_col):
            ws.column_dimensions[get_column_letter(col)].width = 7
        for col, width in zip(
            range(location_start_col, location_start_col + 4),
            (9, 12, 12, 12),
        ):
            ws.column_dimensions[get_column_letter(col)].width = width
        for row in range(1, total_row + 1):
            ws.row_dimensions[row].height = 18
        ws.row_dimensions[1].height = 36
        for row in range(8, 11):
            ws.row_dimensions[row].height = 22
        ws.row_dimensions[11].height = 54
        _align_range(ws, 1, 10, 1, max_col)
        _align_range(ws, 1, 1, 1, max_col, bold=True)
        self._disable_aggregate_info_wrapping(ws)
        _style_range(
            ws,
            6,
            10,
            aggregate_start_col,
            location_start_col - 1,
            fill=LIGHT_BLUE_FILL,
            bold=True,
        )
        _style_range(ws, 11, 11, 1, max_col, fill=LIGHT_BLUE_FILL, bold=True)
        if subtotal_row > data_start:
            _style_range(ws, data_start, subtotal_row - 1, 1, max_col)
        _style_range(ws, subtotal_row, total_row, 1, max_col, fill=LABEL_FILL, bold=True)
        self._format_aggregate_single_line_cells(
            ws,
            data_start,
            total_row,
            location_start_col,
            max_col,
        )
        ws.cell(1, 1).font = Font(bold=True, size=16)

    @staticmethod
    def _disable_aggregate_info_wrapping(ws) -> None:
        for row in (3, 4, 5, 7, 8, 9, 10):
            for col in range(1, 12):
                ws.cell(row, col).alignment = Alignment(
                    horizontal="center", vertical="center", wrap_text=False
                )

    @staticmethod
    def _format_aggregate_single_line_cells(
        ws,
        data_start: int,
        total_row: int,
        location_start_col: int,
        max_col: int,
    ) -> None:
        for row in range(11, total_row + 1):
            ws.cell(row, 1).alignment = Alignment(
                horizontal="center", vertical="center", wrap_text=False
            )
            for col in range(location_start_col, max_col + 1):
                ws.cell(row, col).alignment = Alignment(
                    horizontal="center", vertical="center", wrap_text=False
                )
        for row in range(data_start, total_row):
            ws.cell(row, 1).alignment = Alignment(
                horizontal="center", vertical="center", wrap_text=False
            )

    def _build_aggregate_group_headers(
        self,
        ws,
        aggregate_start_col: int,
        defect_start_col: int,
        condition_start_col: int,
        location_start_col: int,
    ) -> None:
        ws.merge_cells(
            start_row=6,
            start_column=aggregate_start_col,
            end_row=10,
            end_column=defect_start_col - 1,
        )
        ws.cell(6, aggregate_start_col, "합계")
        ws.merge_cells(
            start_row=6,
            start_column=defect_start_col,
            end_row=6,
            end_column=location_start_col - 1,
        )
        ws.cell(6, defect_start_col, "이상항목")
        ws.merge_cells(
            start_row=7,
            start_column=defect_start_col,
            end_row=7,
            end_column=condition_start_col - 1,
        )
        ws.cell(7, defect_start_col, "관로")
        ws.merge_cells(
            start_row=7,
            start_column=condition_start_col,
            end_row=7,
            end_column=location_start_col - 1,
        )
        ws.cell(7, condition_start_col, "상태")
        for col, grade in zip(range(aggregate_start_col, defect_start_col), GRADE_ORDER):
            ws.cell(11, col, grade)

        col = defect_start_col
        for item in AGGREGATE_DEFECT_ITEMS:
            ws.merge_cells(start_row=8, start_column=col, end_row=10, end_column=col + 2)
            label = item.replace("(", "\n(")
            ws.cell(8, col, f"(관로)\n{label}")
            for grade_col, grade in zip(range(col, col + 3), GRADE_ORDER):
                ws.cell(11, grade_col, grade)
            col += 3
        for item in AGGREGATE_CONDITION_ITEMS:
            ws.merge_cells(start_row=8, start_column=col, end_row=10, end_column=col + 2)
            ws.cell(8, col, AGGREGATE_CONDITION_LABELS[item])
            for grade_col, grade in zip(range(col, col + 3), GRADE_ORDER):
                ws.cell(11, grade_col, grade)
            col += 3
        for col, header in zip(
            range(location_start_col, location_start_col + 4),
            ["시·도", "시·군·구", "읍·면·동", "조사일자"],
        ):
            ws.cell(11, col, header)

    def _fill_aggregate_report_row(
        self, ws, row: int, data: ReportWorkbookData
    ) -> None:
        report = data["report"]
        pipe_info = data["pipe_info"]
        upstream = data["upstream_manhole"]
        downstream = data["downstream_manhole"]
        defects = [
            defect
            for defect in data["defects"]
            if (defect.get("item_category") or "관로") == "관로"
        ]
        summary = compute_pipe_state_grades(defects)
        length_m = _float_value(pipe_info, "length_m")
        drive_m = _float_value(pipe_info, "total_drive_distance_m")
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
        ws.cell(row, 1, _value(report, "report_number"))
        fixed_values = [
            _value(report, "pipe_number"),
            _value(upstream, "manhole_number"),
            _value(downstream, "manhole_number"),
            _value(report, "drain_type"),
            _value(report, "pipe_type"),
            _value(report, "specification"),
            _number_or_blank(length_m),
            _number_or_blank(drive_m),
            _number_or_blank(summary.structural_grade),
            _number_or_blank(summary.operational_grade),
        ]
        for col, value in enumerate(fixed_values, start=3):
            ws.cell(row, col, value)

        grade_counts = {grade: 0 for grade in GRADE_ORDER}
        defect_counts = {
            item: {grade: 0 for grade in GRADE_ORDER}
            for item in AGGREGATE_DEFECT_ITEMS
        }
        condition_counts = {item: 0 for item in AGGREGATE_CONDITION_ITEMS}
        for defect in defects:
            defect_item = defect.get("defect_item") or ""
            condition_item = defect.get("condition_item") or ""
            grade = defect.get("grade") or ""
            if defect_item and grade in GRADE_ORDER:
                grade_counts[grade] += 1
                if defect_item in defect_counts:
                    defect_counts[defect_item][grade] += 1
            elif condition_item in condition_counts:
                condition_counts[condition_item] += 1

        col = 13
        for grade in GRADE_ORDER:
            ws.cell(row, col, grade_counts[grade] or None)
            col += 1
        for item in AGGREGATE_DEFECT_ITEMS:
            for grade in GRADE_ORDER:
                ws.cell(row, col, defect_counts[item][grade] or None)
                col += 1
        for item in AGGREGATE_CONDITION_ITEMS:
            ws.cell(row, col, None)
            ws.cell(row, col + 1, condition_counts[item] or None)
            ws.cell(row, col + 2, None)
            col += 3
        location_values = [
            _value(report, "province"),
            _value(report, "city_county"),
            _value(report, "town"),
            _value(report, "survey_date"),
        ]
        location_start_col = 16 + (len(AGGREGATE_DEFECT_ITEMS) * 3) + (
            len(AGGREGATE_CONDITION_ITEMS) * 3
        )
        for location_col, value in zip(
            range(location_start_col, location_start_col + 4), location_values
        ):
            ws.cell(row, location_col, value)

    @staticmethod
    def _report_number_range(reports: list[ReportWorkbookData]) -> str:
        numbers = [str(_value(item["report"], "report_number")) for item in reports]
        numbers = [number for number in numbers if number]
        if not numbers:
            return ""
        return numbers[0] if len(numbers) == 1 else f"{numbers[0]}~{numbers[-1]}"

    def _build_basic_sheet(self, ws, context, report) -> None:
        ws.column_dimensions[get_column_letter(1)].width = 18
        ws.column_dimensions[get_column_letter(2)].width = 28
        ws.column_dimensions[get_column_letter(3)].width = 18
        ws.column_dimensions[get_column_letter(4)].width = 28

        row = 1
        row = _append_title(ws, "프로젝트/사업 요약", row)
        row = _append_pair_table(
            ws,
            row,
            [
                ("프로젝트명", _value(context, "project_name")),
                ("프로젝트 생성일", _value(context, "project_created_at")),
                ("사업코드", _value(context, "business_code")),
                ("사업명", _value(context, "business_name")),
                ("발주처", _value(context, "client")),
                ("사업시작일", _value(context, "business_start_date")),
                ("사업완료일", _value(context, "business_end_date")),
            ],
        )
        row = _append_title(ws, "보고서 기본정보", row)
        row = _append_pair_table(
            ws,
            row,
            [
                ("보고서번호", _value(report, "report_number")),
                ("관로번호", _value(report, "pipe_number")),
                ("조사목적", _value(report, "survey_purpose")),
                ("조사일자", _value(report, "survey_date")),
                ("매설년수(년)", _value(report, "buried_years")),
                ("조사자", _value(report, "inspector")),
                ("시공자", _value(report, "contractor")),
            ],
        )
        row = _append_title(ws, "위치/배수 정보", row)
        row = _append_pair_table(
            ws,
            row,
            [
                ("처리구역", _value(report, "treatment_area")),
                ("배수구역", _value(report, "drainage_area")),
                ("배수분구", _value(report, "drainage_district")),
                ("배수방식", _value(report, "drain_type")),
                ("배수체계", _value(report, "drain_system")),
                ("시도", _value(report, "province")),
                ("시군구", _value(report, "city_county")),
                ("읍면동", _value(report, "town")),
                ("리", _value(report, "village")),
                ("지번", _value(report, "lot_number")),
                ("도로명주소", _value(report, "road_address")),
            ],
        )
        row = _append_title(ws, "관로 분류 정보", row)
        row = _append_pair_table(
            ws,
            row,
            [
                ("관종", _value(report, "pipe_type")),
                ("구분", _value(report, "category")),
                ("규격", _value(report, "specification")),
                ("보고서 생성일", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
            ],
        )

    def _build_pipe_sheet(
        self, ws, pipe_info, upstream_manhole, downstream_manhole, actual_survey
    ) -> None:
        headers = [
            "구분",
            "맨홀번호",
            "맨홀종류",
            "맨홀 내부재질",
            "맨홀 뚜껑재질",
            "맨홀 E.L(m)",
            "맨홀크기",
            "맨홀깊이(m)",
            "맨홀 인버트",
            "사다리모양",
            "위도",
            "경도",
        ]
        for idx, header in enumerate(headers, start=1):
            ws.cell(row=1, column=idx, value=header)
            ws.cell(row=1, column=idx).font = Font(bold=True)
            ws.cell(row=1, column=idx).fill = HEADER_FILL
            ws.column_dimensions[get_column_letter(idx)].width = 16

        for row_index, label, manhole in (
            (2, "상류맨홀", upstream_manhole),
            (3, "하류맨홀", downstream_manhole),
        ):
            values = [
                label,
                _value(manhole, "manhole_number"),
                _value(manhole, "manhole_type"),
                _value(manhole, "internal_material"),
                _value(manhole, "cover_material"),
                _value(manhole, "el_m"),
                _value(manhole, "size"),
                _value(manhole, "depth_m"),
                _value(manhole, "invert"),
                _value(manhole, "ladder_shape"),
                _value(manhole, "latitude"),
                _value(manhole, "longitude"),
            ]
            for col, value in enumerate(values, start=1):
                ws.cell(row=row_index, column=col, value=value)

        completed = "완주" if int(_value(pipe_info, "is_completed", 0) or 0) else "미완주"
        pipe_summary = (
            f"관로정보: 연장(m) {_value(pipe_info, 'length_m')} / "
            f"총주행거리(m) {_value(pipe_info, 'total_drive_distance_m')} / "
            f"완주여부 {completed} / "
            f"미주행거리(m) {_value(pipe_info, 'undriven_distance_m')}"
        )
        ws.cell(row=4, column=1, value=pipe_summary)
        ws.merge_cells(start_row=4, start_column=1, end_row=4, end_column=len(headers))
        ws.cell(row=4, column=1).fill = LABEL_FILL
        ws.cell(row=4, column=1).font = Font(bold=True)

        row = 6
        survey_headers = [
            "주행방향->맨홀번호",
            "발생지점(m)",
            "미주행사유",
            "미주행사유설명",
        ]
        for col, header in enumerate(survey_headers, start=1):
            ws.cell(row=row, column=col, value=header)
            ws.cell(row=row, column=col).font = Font(bold=True)
            ws.cell(row=row, column=col).fill = HEADER_FILL

        upstream_no = _value(upstream_manhole, "manhole_number")
        downstream_no = _value(downstream_manhole, "manhole_number")
        rows = [
            [
                f"시작->끝 / {upstream_no} -> {downstream_no}",
                _value(actual_survey, "start_occurrence_point_m"),
                _value(actual_survey, "start_undriven_reason", "없음"),
                _value(actual_survey, "start_undriven_reason_detail"),
            ],
            [
                f"끝->시작 / {downstream_no} -> {upstream_no}",
                _value(actual_survey, "end_occurrence_point_m"),
                _value(actual_survey, "end_undriven_reason", "없음"),
                _value(actual_survey, "end_undriven_reason_detail"),
            ],
        ]
        for offset, values in enumerate(rows, start=1):
            for col, value in enumerate(values, start=1):
                ws.cell(row=row + offset, column=col, value=value)

        content_row = row + 3
        ws.cell(
            row=content_row,
            column=1,
            value=f"조사내용: {_value(actual_survey, 'survey_content')}",
        )
        ws.merge_cells(start_row=content_row, start_column=1, end_row=content_row, end_column=4)
        ws.cell(row=content_row, column=1).alignment = Alignment(wrap_text=True)

    def _build_defect_sheet(
        self, ws, defects: list[DefectReportRow]
    ) -> None:
        headers = [
            "No",
            "시각",
            "주행방향",
            "거리(m)",
            "항목구분",
            "상태항목",
            "이상항목",
            "등급",
            "사분면",
            "맨홀결함깊이(m)",
            "메모",
            "이미지",
        ]
        widths = [8, 12, 12, 10, 12, 20, 20, 8, 16, 16, 28, 22]
        for idx, header in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=idx, value=header)
            cell.font = Font(bold=True)
            cell.fill = HEADER_FILL
            ws.column_dimensions[get_column_letter(idx)].width = widths[idx - 1]

        for index, defect in enumerate(defects, start=1):
            row_index = index + 1
            grade = defect.get("grade") or ""
            category = defect.get("item_category") or ""
            values = [
                index,
                _format_hms(int(defect["timestamp_ms"])),
                defect.get("drive_direction") or "",
                defect.get("distance_m") if defect.get("distance_m") is not None else "",
                category,
                display_condition_item(category, defect.get("condition_item")),
                display_defect_item(category, defect.get("defect_item")),
                grade,
                defect.get("quadrant") or "",
                defect.get("manhole_defect_depth_m")
                if defect.get("manhole_defect_depth_m") is not None
                else "",
                defect.get("memo") or "",
                "",
            ]
            for col, value in enumerate(values, start=1):
                ws.cell(row=row_index, column=col, value=value)
            if grade in GRADE_FILL:
                ws.cell(row=row_index, column=8).fill = GRADE_FILL[grade]

            image_path = Path(defect["image_path"])
            if image_path.exists():
                image = XLImage(str(image_path))
                image.width = 120
                image.height = 68
                ws.add_image(image, f"L{row_index}")
                ws.row_dimensions[row_index].height = 54

        self._append_state_grade_table(ws, len(defects) + 4, defects)

    def _append_state_grade_table(
        self, ws, row_index: int, defects: list[DefectReportRow]
    ) -> None:
        summary = compute_pipe_state_grades(defects)
        ws.cell(row=row_index, column=1, value="단위구간 상태등급")
        ws.merge_cells(start_row=row_index, start_column=1, end_row=row_index, end_column=9)
        title_cell = ws.cell(row=row_index, column=1)
        title_cell.font = Font(bold=True, size=13)
        title_cell.fill = HEADER_FILL
        title_cell.alignment = Alignment(horizontal="center")

        row_index += 1
        ws.cell(row=row_index, column=1, value="전체 구조등급")
        ws.cell(row=row_index, column=2, value=format_state_value(summary.structural_grade))
        ws.cell(row=row_index, column=3, value="전체 운영등급")
        ws.cell(row=row_index, column=4, value=format_state_value(summary.operational_grade))
        for col in (1, 3):
            ws.cell(row=row_index, column=col).fill = LABEL_FILL
            ws.cell(row=row_index, column=col).font = Font(bold=True)

        row_index += 2
        headers = [
            "구간",
            "구간정보",
            "주행방향",
            "거리범위(m)",
            "구조점수",
            "구조등급",
            "운영점수",
            "운영등급",
            "결함수",
        ]
        for col, header in enumerate(headers, start=1):
            cell = ws.cell(row=row_index, column=col, value=header)
            cell.font = Font(bold=True)
            cell.fill = HEADER_FILL

        for offset, section in enumerate(summary.sections, start=1):
            time_range = format_state_time_range(
                section.start_timestamp_ms, section.end_timestamp_ms
            )
            section_label = format_state_section_label(section)
            section_info = " / ".join(
                part for part in (time_range, section_label) if part
            )
            distance_range = (
                f"{format_state_distance(section.start_distance_m)}~"
                f"{format_state_distance(section.end_distance_m)}"
            )
            values = [
                section.index,
                section_info,
                section.drive_direction,
                distance_range,
                section.structural_score,
                section.structural_grade,
                section.operational_score,
                section.operational_grade,
                section.defect_count,
            ]
            for col, value in enumerate(values, start=1):
                ws.cell(row=row_index + offset, column=col, value=value)
