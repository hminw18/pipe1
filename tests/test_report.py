from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from sewerpipe_inspector.db import Database
from sewerpipe_inspector.services.inspection_service import InspectionService
from sewerpipe_inspector.services.report_service import (
    AGGREGATE_CONDITION_ITEMS,
    AGGREGATE_DEFECT_ITEMS,
)
from sewerpipe_inspector.services.report_service import ReportService
from sewerpipe_inspector.services.storage_service import StorageService


def test_generate_inspection_report_creates_multi_sheet_workbook(tmp_path: Path) -> None:
    capture = tmp_path / "capture.png"
    capture.write_bytes(
        b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01\x00\x00\x00\x01\x08\x02\x00\x00\x00\x90wS\xde\x00\x00\x00\x0cIDAT\x08\x1dc\xf8\xcf\xc0\x00\x00\x04\x01\x01\x00\x18\xdd\x8d\x17\x00\x00\x00\x00IEND\xaeB`\x82"
    )

    db = Database(tmp_path / "app.db")
    storage = StorageService(tmp_path / "workspace")
    inspection = InspectionService(db, storage, ReportService())

    project_id = db.create_project("Project-1")
    business_id = db.create_business(
        project_id, "B001", "Business-1", "Client", "2026-01-01", "2026-12-31"
    )
    report_id = db.create_report(business_id, "R001", "PIPE-001")
    report = db.get_report(report_id)
    payload = {field: report[field] for field in report.keys()}
    payload.update({"survey_date": "2026-06-13", "buried_years": "10"})
    db.update_report(report_id, payload)
    db.update_pipe_information(report_id, 20.0, 18.0)
    db.update_manhole(report_id, "upstream", {"manhole_number": "U-1"})
    db.update_manhole(report_id, "downstream", {"manhole_number": "D-1"})
    db.update_actual_survey(
        report_id,
        {
            "start_occurrence_point_m": 1.0,
            "start_undriven_reason": "없음",
            "start_undriven_reason_detail": None,
            "end_occurrence_point_m": 2.0,
            "end_undriven_reason": "기타",
            "end_undriven_reason_detail": "memo",
            "survey_content": "조사내용",
        },
    )
    video_id = db.upsert_video(report_id, str(tmp_path / "a.mp4"), 10.0, None, "순주행")
    db.create_defect(
        report_id, video_id, 1200, str(capture), "순주행", 1.25,
        "관로", "구조특징", "균열(길이)", "대", "상", None, "minor"
    )

    output = inspection.generate_excel_report(report_id)

    assert output.exists()
    workbook = load_workbook(output)
    assert workbook.sheetnames == ["기본정보", "관로조사정보", "결함목록"]
    assert workbook["기본정보"]["A1"].value == "프로젝트/사업 요약"
    assert "관로정보:" in workbook["관로조사정보"]["A4"].value
    assert workbook["결함목록"]["F2"].value == "구조특징"
    assert workbook["결함목록"]["G2"].value == "균열(길이) (CL)"
    assert workbook["결함목록"]["H2"].value == "대"


def test_generate_business_excel_workbooks_for_selected_reports(tmp_path: Path) -> None:
    capture = tmp_path / "capture.png"
    capture.write_bytes(
        b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01\x00\x00\x00\x01\x08\x02\x00\x00\x00\x90wS\xde\x00\x00\x00\x0cIDAT\x08\x1dc\xf8\xcf\xc0\x00\x00\x04\x01\x01\x00\x18\xdd\x8d\x17\x00\x00\x00\x00IEND\xaeB`\x82"
    )

    db = Database(tmp_path / "app.db")
    storage = StorageService(tmp_path / "workspace")
    inspection = InspectionService(db, storage, ReportService())

    project_id = db.create_project("Project-1")
    business_id = db.create_business(
        project_id, "B001", "Business-1", "Client", "2026-01-01", "2026-12-31"
    )
    report_ids: list[int] = []
    for index in range(2):
        report_id = db.create_report(business_id, f"R{index + 1:03d}", f"PIPE-{index + 1:03d}")
        report = db.get_report(report_id)
        payload = {field: report[field] for field in report.keys()}
        payload.update(
            {
                "survey_date": "2026-06-13",
                "pipe_type": "흄관_HP",
                "specification": "450",
                "drain_type": "오수",
                "province": "경기",
                "city_county": "의정부시",
                "town": "가능동",
            }
        )
        db.update_report(report_id, payload)
        db.update_pipe_information(report_id, 20.0, 18.0)
        db.update_manhole(report_id, "upstream", {"manhole_number": "MH01"})
        db.update_manhole(report_id, "downstream", {"manhole_number": "MH02"})
        video_id = db.upsert_video(
            report_id, str(tmp_path / f"{index}.mp4"), 10.0, None, "순주행"
        )
        db.create_defect(
            report_id,
            video_id,
            0,
            str(capture),
            "순주행",
            0.0,
            "관로",
            "조사시작(순방향)",
            None,
            None,
            None,
            None,
            "start",
        )
        db.create_defect(
            report_id,
            video_id,
            1000,
            str(capture),
            "순주행",
            1.0,
            "관로",
            None,
            "균열(길이)",
            "대",
            "상",
            None,
            "",
        )
        db.create_defect(
            report_id,
            video_id,
            2000,
            str(capture),
            "순주행",
            2.0,
            "관로",
            None,
            "토사퇴적",
            "중",
            "상",
            None,
            "",
        )
        db.create_defect(
            report_id,
            video_id,
            3000,
            str(capture),
            "순주행",
            3.0,
            "관로",
            "조사중단",
            None,
            None,
            None,
            None,
            "stop",
        )
        report_ids.append(report_id)

    internal_path = inspection.generate_internal_defect_workbook(
        report_ids, tmp_path / "internal.xlsx"
    )
    aggregate_path = inspection.generate_defect_aggregate_workbook(
        report_ids, tmp_path / "aggregate.xlsx"
    )

    internal = load_workbook(internal_path, data_only=False)
    assert internal.sheetnames == [
        "내부조사(CCTV)집계표",
        "R001",
        "R002",
    ]
    assert internal["내부조사(CCTV)집계표"]["C2"].value == "하수관거 단위구간 보고서"
    summary = internal["내부조사(CCTV)집계표"]
    assert summary["E4"].value == "시점맨홀번호"
    assert summary["F4"].value == "종점맨홀번호"
    assert summary["I4"].value == "관로연장"
    assert summary["J4"].value == "조사연장"
    assert summary["AB3"].value == "단위구간(개소)"
    assert summary["AB4"].value == "단위구간(개소)"
    assert summary.column_dimensions["E"].width == 16
    assert summary.column_dimensions["F"].width == 16
    assert summary.column_dimensions["AB"].width == 16
    assert len(internal["내부조사(CCTV)집계표"]._images) == 0
    assert internal["R001"]["A1"].value == "<내부결함 판독표>"
    assert "A1:L1" in [str(merged) for merged in internal["R001"].merged_cells.ranges]
    assert internal["R001"]["A1"].fill.fill_type is None
    assert internal["R001"]["M4"].fill.fill_type is None
    assert internal["R001"].column_dimensions["M"].width == 4.5
    assert internal["R001"]["D10"].value == "균열(길이)(대)"
    assert internal["R001"]["F10"].value == 40
    assert internal["R001"]["J9"].value.startswith("=IF(I9>69")
    assert len(internal["R001"]._images) == 0

    internal_with_photos_path = inspection.generate_internal_defect_workbook(
        report_ids, tmp_path / "internal_with_photos.xlsx", include_photos=True
    )
    internal_with_photos = load_workbook(internal_with_photos_path, data_only=False)
    assert internal_with_photos.sheetnames == [
        "내부조사(CCTV)집계표",
        "R001",
        "R002",
    ]
    assert len(internal_with_photos["내부조사(CCTV)집계표"]._images) == 0
    assert internal_with_photos["R001"]["M6"].value == "사진"
    assert len(internal_with_photos["R001"]._images) == 4

    aggregate = load_workbook(aggregate_path, data_only=False)
    sheet = aggregate.active
    aggregate_start_col = 13
    defect_start_col = aggregate_start_col + 3
    condition_start_col = defect_start_col + (len(AGGREGATE_DEFECT_ITEMS) * 3)
    location_start_col = condition_start_col + (len(AGGREGATE_CONDITION_ITEMS) * 3)
    max_col = location_start_col + 3
    max_letter = get_column_letter(max_col)
    assert sheet["A1"].value == "하수관거 조사 집계표"
    assert sheet["C7"].value == "2건"
    assert f"C3:{max_letter}3" in [str(merged) for merged in sheet.merged_cells.ranges]
    assert f"C4:{max_letter}4" in [str(merged) for merged in sheet.merged_cells.ranges]
    assert sheet["C3"].alignment.wrap_text is not True
    assert sheet["A4"].alignment.wrap_text is not True
    assert sheet["A3"].fill.fill_type is None
    assert sheet["A3"].border.left.style is None
    assert sheet["A11"].font.bold
    assert sheet["A11"].fill.fill_type == "solid"
    assert "A11:B11" in [str(merged) for merged in sheet.merged_cells.ranges]
    assert "A12:B12" in [str(merged) for merged in sheet.merged_cells.ranges]
    assert sheet["A12"].alignment.wrap_text is not True
    assert sheet.row_dimensions[8].height == 22
    assert sheet.row_dimensions[11].height == 54
    assert sheet.cell(6, aggregate_start_col).font.bold
    assert sheet.cell(6, aggregate_start_col).fill.fill_type == "solid"
    assert sheet["A12"].value == "R001"
    assert sheet.cell(12, aggregate_start_col).value == 1
    assert sheet.cell(12, aggregate_start_col + 1).value == 1
    assert sheet.cell(12, defect_start_col).value == 1
    start_condition_col = condition_start_col + (
        AGGREGATE_CONDITION_ITEMS.index("조사시작(순방향)") * 3
    )
    joint_condition_col = condition_start_col + (
        AGGREGATE_CONDITION_ITEMS.index("이음부(접합부)존재") * 3
    )
    assert (
        sheet.cell(8, start_condition_col).value
        == "(상태)\n관로조사시작\n(순방향)"
    )
    assert (
        sheet.cell(8, joint_condition_col).value
        == "(상태)\n이음부(접합부)존재"
    )
    assert sheet.column_dimensions[get_column_letter(start_condition_col)].width == 7
    assert sheet.cell(11, location_start_col).value == "시·도"
    assert sheet.cell(11, location_start_col).alignment.wrap_text is not True
    assert sheet.cell(12, location_start_col).value == "경기"
    assert sheet.cell(12, location_start_col).alignment.wrap_text is not True
    assert sheet.column_dimensions[get_column_letter(location_start_col)].width == 9
    assert sheet.column_dimensions[get_column_letter(location_start_col + 3)].width == 12
    silt_col = defect_start_col + (AGGREGATE_DEFECT_ITEMS.index("토사퇴적") * 3)
    stop_col = condition_start_col + (AGGREGATE_CONDITION_ITEMS.index("조사중단") * 3)
    assert sheet[f"{get_column_letter(silt_col + 1)}12"].value == 1
    assert sheet[f"{get_column_letter(stop_col + 1)}12"].value == 1
